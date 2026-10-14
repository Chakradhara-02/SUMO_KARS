"""Build the short explanation of the current SUMO KARS reachable-set pipeline."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / 'docs' / 'Reachable_Set_and_Current_Code_Explanation.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.65)
sec.left_margin = sec.right_margin = Inches(.75)
sec.footer_distance = Inches(.28)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.widow_control = True
    if style._element.pPr is not None:
        for border in list(style._element.pPr.findall(qn('w:pBdr'))):
            style._element.pPr.remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Normal'].paragraph_format.space_after = Pt(4)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Title'].paragraph_format.space_after = Pt(10)
for name, size in [('Heading 1',15),('Heading 2',12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(5)
    doc.styles[name].paragraph_format.keep_with_next = True


def para(text, bold_lead=None):
    after_table = len(doc._element.body) > 1 and doc._element.body[-2].tag == qn('w:tbl')
    p=doc.add_paragraph()
    if after_table: p.paragraph_format.space_before=Pt(6)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold=True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    for row in rows:
        cells=t.add_row().cells
        for cell,text in zip(cells,row):cell.text=text
    borders=OxmlElement('w:tblBorders')
    for tag in ('top','left','bottom','right','insideH','insideV'):
        border=OxmlElement('w:'+tag)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:border.set(qn('w:'+k),v)
        borders.append(border)
    t._tbl.tblPr.append(borders)
    for index,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        if index==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                elem=OxmlElement('w:'+side);elem.set(qn('w:w'),'80');elem.set(qn('w:type'),'dxa');margins.append(elem)
            props.append(margins)
            if index==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EEF2');props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.line_spacing=1.03
                for r in p.runs:r.font.size=Pt(10.5);r.bold=index==0
    return t


def mr(text):
    r=OxmlElement('m:r');text_node=OxmlElement('m:t');text_node.text=text;r.append(text_node);return r


def sub(base,index):
    node=OxmlElement('m:sSub')
    e=OxmlElement('m:e');e.append(mr(base));node.append(e)
    s=OxmlElement('m:sub');s.append(mr(index));node.append(s)
    return node


def equation(kind):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(4)
    math_para=OxmlElement('m:oMathPara');math=OxmlElement('m:oMath')
    math.append(sub(kind,'next'))
    math.append(mr(' = '+kind))
    if kind=='v':math.append(mr(' + a Δt'))
    else:
        math.append(mr(' + v Δt + '))
        frac=OxmlElement('m:f')
        for tag,text in [('num','1'),('den','2')]:
            n=OxmlElement('m:'+tag);n.append(mr(text));frac.append(n)
        math.append(frac);math.append(mr(' a '))
        power=OxmlElement('m:sSup');base=OxmlElement('m:e');base.append(mr('Δt'));power.append(base)
        exponent=OxmlElement('m:sup');exponent.append(mr('2'));power.append(exponent);math.append(power)
    math_para.append(math);p._p.append(math_para)


footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
field_run=OxmlElement('w:r')
field_text=OxmlElement('w:t');field_text.text='1';field_run.append(field_text)
field.append(field_run);footer._p.append(field)

# Page one introduces the concept and identifies the actual data contract.
doc.add_paragraph('Reachable Sets in SUMO KARS',style='Title')
para('The reachable set is the collection of future vehicle states that can be produced from the current state using the allowed driving actions. In this project, it supplies the alternatives from which the planner chooses a lane and speed. This document explains the inputs, the generation and filtering process, and the way the current Python code turns the result into SUMO commands.')
doc.add_heading('What a reachable set means',level=1)
para('A vehicle cannot immediately appear at any position or speed. Its acceleration, braking limits and lane-change duration restrict its future movement. A reachable set records those possibilities over a prediction horizon. The raw set contains dynamically valid alternatives; the adjusted set retains only the paths that pass the implemented traffic-constraint checks. A complete theoretical reachable set may contain continuously many states. This implementation samples actions and keeps a bounded graph.')
para('Each State stores longitudinal position x in metres, speed v in metres per second, lane index, absolute simulation time t, lane_hold and lane_origin. The last two fields record the remaining lane-change time and the original lane. Nodes represent states; directed edges represent allowed transitions between consecutive prediction times.')
doc.add_heading('What data enters the planner',level=1)
table(['Input','Where it comes from','How it is used'],[
    ('Ego state','TraCI position, speed, lane and time','Starting node of the reachable-set graph'),
    ('Other vehicles','Current local observations and delivered IoT snapshots','Predicted occupied road regions, using actual vehicle dimensions'),
    ('Road restrictions','Map rules, road events and signal phase information','Lane, position, speed and time constraints'),
    ('Configuration','config.yaml and run overrides','Dynamics, horizon, sampling limits and cost weights'),
], [1.15,2.55,3.3])
para('Traffic is sampled every 0.2 s. The default local sensing range is 120 m; current nearby observations replace older remote observations. IoT packets carry vehicle records, constraints and their observation time. The Channel simulates delay, packet loss, jitter and outages and rejects older arrivals that would overwrite a newer snapshot.')
doc.add_heading('Current default settings',level=2)
para('Planning interval 1 s; prediction horizon 5 steps or 5 s; speed 0–16 m/s; acceleration actions 0, +2, −2 and −4 m/s²; three main-road lanes; lane-change duration 2.5 s. Position and speed bins are 2 m and 1 m/s, with at most 180 retained nodes per future layer.')

para('Code locations: kars/models.py, kars/reachable_set.py, kars/knowledge.py, kars/planning.py, kars/control.py, kars/simulation.py and config.yaml. Paths are relative to SUMO_KARS.').runs[0].font.size=Pt(9)

# Page two explains the graph algorithms and gives a checked numerical example.
doc.add_page_break()
doc.add_heading('How the reachable set is generated',level=1)
para('1  Prepare driving actions. offline_rs() produces an acceleration list and a transition table for integer speeds within the configured range. Planner prepares this once at initialization. The online generator uses the action list and recomputes transitions from the actual measured state; it does not simply look up an integer-speed table entry.',bold_lead='1  Prepare driving actions.')
para('2  Expand future states. online_rs() creates a NetworkX directed graph with the current state as node 0. For each of the five future layers, it calls propagate() for each sampled acceleration and each permitted lane: the current lane or an adjacent lane. The constant-acceleration equations are:',bold_lead='2  Expand future states.')
equation('v');equation('x')
para('Here a is acceleration and Δt is the 1 s planning interval. propagate() rejects speed-limit violations, invalid lane indices, changes across more than one lane and a new lane change while lane_hold is positive. During a lane change, constraint checks include both the origin and destination lanes.')
para('Numerical example. Starting at x = 20 m, v = 8 m/s, lane 1 and t = 10 s, choosing a = +2 m/s² while staying in lane 1 produces x = 29 m, v = 10 m/s and t = 11 s. Choosing a = −4 m/s² instead produces x = 26 m and v = 4 m/s. These are two possible successors, not a statement of which action the planner will select.',bold_lead='Numerical example.')
para('3  Bound graph size. States are assigned position and velocity bins, but nodes are merged only when their continuous positions and speeds match. A different state colliding with an existing bin is discarded. If a layer exceeds 180 nodes, the code allocates samples across lanes and spreads them over the sorted states. Sampling and truncation make the graph an approximation.',bold_lead='3  Bound graph size.')
doc.add_heading('How traffic knowledge changes the set',level=1)
para('KnowledgeGraph.update() writes fused observations and constraints as RDF triples. For kars and kars_delay, constraints() executes knowledge_graph/queries/constraints.rq and returns Python Constraint objects. vehicle_constraints() predicts main-road vehicle occupancy. The age-aware version spans braking and constant-speed predictions and adds 2 m of margin per second of observation age to the base 8 m safety gap.')
table(['Constraint kind','Meaning of a rejected transition'],[
    ('spatial','The swept position interval intersects a blocked region in an occupied lane.'),
    ('spatial_speed','The interval intersects a speed zone and either endpoint speed exceeds its cap.'),
    ('spatial_temporal','A blocked region overlaps both the movement interval and its active time.'),
    ('spatial_temporal_speed','An active time-dependent speed restriction is exceeded.'),
    ('marking','A lane change crosses the restricted marking region.'),
], [2.0,5.0])
para('rs_adjustment() removes violating edges, records their reasons and deletes nodes that can no longer be reached from node 0. Its outputs are the adjusted graph and rejected-transition records. A single surviving state is insufficient: the planner needs a connected path reaching the final horizon layer.')

# Page three connects the result to optimization, vehicle commands and files.
doc.add_page_break()
doc.add_heading('How the simulation uses the output',level=1)
doc.add_heading('Selecting a corridor and trajectory',level=2)
para('corridors() groups final-layer states by ending lane and collects their ancestor paths. Candidates may share states. Their margin proxy counts distinct sampled positions per layer; it is not physical clearance in metres.')
para('dynamic_programming() minimizes accumulated cost in each candidate. Cost penalizes distance from the centre lane, lower speed, acceleration and lane changes. Candidate score equals path cost minus 0.05 times the margin proxy; the lowest score wins. astar uses A* on the constrained graph; rs keeps occupancy checks but omits speed and marking rules.')
para('optimize_trajectory() uses SciPy SLSQP to smooth longitudinal accelerations within acceleration, speed and position bounds. It penalizes acceleration, acceleration changes and position deviation. Every proposed segment is rechecked against the constraints. Failed smoothing falls back to the original path. Lanes remain discrete intents; world_y describes intended gradual lateral movement.')
doc.add_heading('The current closed loop',level=2)
para('run_simulation() advances SUMO, collects traffic, applies events, delivers IoT packets and fuses observations. It passes the current State and constraints to Planner.plan() roughly every 1 s or when delivered infrastructure restrictions change. Periodic replanning also responds to vehicle movement.')
para('The plan returns raw and adjusted graphs, rejection reasons, candidates, selected_lane, chain, trajectory, emergency and timings. The trajectory contains time, x, y, world_y, velocity, acceleration, jerk and lane arrays, plus an optimized flag. Only velocity[1] and lane[1] become the next targets; the full prediction is not executed at once.')
para('TrajectoryTracker.command() bounds speed requests every 0.2 s and commits an in-progress lane change until completion or timeout. TraCI setSpeed() and changeLane() send requests. SUMO executes gradual movement with braking, collision, signal and right-of-way safeguards enabled.')
para('If no full-horizon path survives, emergency braking is requested. An imminent all-lane blockage also overrides the plan. kars_delay additionally brakes for missing or more than 3 s old roadside data. At the default −4 m/s² limit, emergency() reduces requested speed by up to 0.8 m/s per step. Replanning permits recovery when conditions allow.')
doc.add_heading('What is saved and displayed',level=2)
table(['Output','What it contains'],[
    ('live.json and plans.json','Reachable states, surviving edges, selected chain, corridor scores and planning snapshots'),
    ('trajectory.csv','The last available feasible planned trajectory, not the entire executed trip'),
    ('ego.csv and traffic.csv','Measured SUMO movement, vehicle dimensions and tracking data'),
    ('trace.jsonl and metrics.json','Recorded replay frames and measured safety, efficiency, comfort and timing results'),
], [2.05,4.95])
para('Implementation scope. This sampled lane-level graph can omit feasible paths. Main-road predictions use simplified motion models; crossing conflicts also rely on SUMO safety. Feasibility is not a formal safety guarantee. Compare predictions with measured execution.',bold_lead='Implementation scope.')


doc.core_properties.title='Reachable Sets in SUMO KARS'
doc.core_properties.subject='Reachable-set inputs, graph generation, constraint adjustment and closed-loop execution'
doc.core_properties.author='SUMO KARS'
doc.core_properties.keywords='reachable set, SUMO, KARS, TraCI, SPARQL, dynamic programming'
doc.save(DEST)
print(DEST)
print('Words:',sum(len(p.text.split()) for p in doc.paragraphs)+sum(len(c.text.split()) for t in doc.tables for r in t.rows for c in r.cells))
